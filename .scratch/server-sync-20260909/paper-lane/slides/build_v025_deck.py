#!/usr/bin/env python3
"""Author the V0.25 physics-successor teaching-deck skeleton on the wmnlab template.

Lineage
-------
This is `pilot-build/build_pilot.py` (V0.23) generalised from one pilot slide to
the full V0.25 arc.  It keeps that script's contract verbatim:

* the wmnlab template supplies the master, logo, rules and slide-number field;
* every slide is authored as plain shapes on a blank layout;
* mathematics is NOT drawn as text -- each equation is a named autoshape whose
  only paragraph is the marker `[[PPT_NATIVE_MATH:<id>]]`, replaced in a second
  pass by a native `a14:m` / OMML math zone (see `native_math.py`);
* typography is Times New Roman at exactly three sizes: 28 pt slide titles,
  24 pt ordinary body text, 20 pt text inside cards / boxes / badges.

Two deviations from the V0.23 script are forced by this machine and are
recorded in BUILD-NOTES.md:
  1. the template path `/home/u24/pptx-craft/assets/wmnlab.pptx` does not exist
     here; the same asset is resolved at `/home/sat/pptx-wrap/assets/templates/`;
  2. the pptx-craft OMML injector that consumed `native-formulas.json` is not
     installed; `native_math.py` reimplements that stage against the OMML
     dialect recovered from `pilot-build/lc-srs-pilot-native-v4.pptx`.

No result numbers appear anywhere in the deck.  Every place a result belongs
carries the placeholder `⟨結果待填⟩`.
"""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import native_math

HERE = Path(__file__).resolve().parent
TEMPLATE = Path("/home/sat/pptx-wrap/assets/templates/wmnlab.pptx")
MANIFEST = HERE / "native-formulas-v025.json"
OUTPUT = HERE / "v025-deck-skeleton.pptx"
REPORT = HERE / "build-report.json"

FONT = "Times New Roman"
TITLE_PT, BODY_PT, BOX_PT = 28, 24, 20
PENDING = "⟨結果待填⟩"

BLUE = RGBColor(0x32, 0x32, 0x83)
DEEP_BLUE = RGBColor(0x00, 0x33, 0x66)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
PANEL = RGBColor(0xF8, 0xFA, 0xFC)
PALE_BLUE = RGBColor(0xF0, 0xF4, 0xF8)
PALE_GREEN = RGBColor(0xE8, 0xF5, 0xE9)
PALE_RED = RGBColor(0xFE, 0xF2, 0xF2)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xFF, 0xF3, 0xCD)
AMBER_LINE = RGBColor(0xD9, 0x77, 0x06)
AMBER_TEXT = RGBColor(0x85, 0x64, 0x04)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BADGES = {
    "SUCCESSION": (AMBER, AMBER_LINE, AMBER_TEXT),
    "FROZEN": (PALE_GREEN, GREEN, RGBColor(0x14, 0x53, 0x2D)),
    "RESULTS_PENDING": (PALE_RED, RED, RGBColor(0x7F, 0x1D, 0x1D)),
    "CARRIED": (PALE_BLUE, BLUE, DEEP_BLUE),
}


# --------------------------------------------------------------------------
# primitives (lifted from build_pilot.py)
# --------------------------------------------------------------------------

def _remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def _set_text(shape, text, *, size, bold=False, color=INK,
              align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.08):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = color


def _set_lines(shape, lines, *, size, color=INK, margin=0.10,
               valign=MSO_ANCHOR.TOP, space_after=6):
    """lines: sequence of (text, bold, indent_level)."""
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, item in enumerate(lines):
        text, bold, level = item
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.level = min(level, 4)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = False
        run.font.color.rgb = DEEP_BLUE if bold else color


def _text_box(slide, x, y, w, h, text, *, size=BOX_PT, bold=False, color=INK,
              align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _set_text(shape, text, size=size, bold=bold, color=color, align=align)
    return shape


def _card(slide, x, y, w, h, *, fill=PANEL, line=LINE, radius=True):
    kind = (MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius
            else MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def _formula_host(slide, formula_id, x, y, w, h, *, fill=WHITE, line=LINE):
    shape = _card(slide, x, y, w, h, fill=fill, line=line)
    shape.name = f"Math host {formula_id}"
    shape.element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}nvSpPr"
    ).find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
    ).set("descr", f"Native Office Math: {formula_id}")
    _set_text(shape, f"[[PPT_NATIVE_MATH:{formula_id}]]", size=BOX_PT,
              color=INK, align=PP_ALIGN.CENTER, margin=0.04)
    return shape


# --------------------------------------------------------------------------
# slide chrome
# --------------------------------------------------------------------------

def _chrome(slide, title, badge, subtitle):
    shape = _text_box(slide, 0.62, 0.13, 9.05, 0.66, title, size=TITLE_PT,
                      bold=True, color=BLUE)
    shape.name = "Slide title"
    if badge:
        fill, line, text_color = BADGES[badge]
        chip = _card(slide, 9.76, 0.20, 2.95, 0.44, fill=fill, line=line)
        chip.name = "Status badge"
        _set_text(chip, badge, size=BOX_PT, bold=True, color=text_color,
                  align=PP_ALIGN.CENTER, margin=0.02)
    if subtitle:
        shape = _text_box(slide, 0.64, 0.85, 11.95, 0.56, subtitle,
                          size=BODY_PT, color=INK)
        shape.name = "Slide subtitle"


def _footer(slide, text, *, tone="amber"):
    fill, line = (AMBER, AMBER_LINE) if tone == "amber" else (PALE_BLUE, BLUE)
    band = _card(slide, 0.62, 6.46, 12.09, 0.88, fill=fill, line=line)
    band.name = "Footer band"
    _set_text(band, text, size=BOX_PT, color=INK, align=PP_ALIGN.CENTER,
              margin=0.06)


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# --------------------------------------------------------------------------
# slide bodies
# --------------------------------------------------------------------------

def render_title(slide, spec):
    _card(slide, 0.62, 1.55, 12.09, 2.55, fill=PALE_BLUE, line=BLUE)
    _text_box(slide, 0.95, 1.75, 11.4, 0.90, spec["title"], size=TITLE_PT,
              bold=True, color=BLUE)
    _text_box(slide, 0.95, 2.66, 11.4, 1.34, spec["subtitle"], size=BODY_PT,
              color=INK)
    box = _card(slide, 0.62, 4.25, 12.09, 2.65, fill=PANEL, line=LINE)
    _set_lines(box, [(line, False, 0) for line in spec["lines"]],
               size=BOX_PT, margin=0.16)


def render_cards(slide, spec):
    cards = spec["cards"]
    columns = len(cards)
    top = 1.45
    height = 4.95 if spec.get("footer") else 5.45
    gap = 0.21
    total = 12.09
    width = (total - gap * (columns - 1)) / columns
    for index, card in enumerate(cards):
        x = 0.62 + index * (width + gap)
        raw = card.get("fill")
        if raw is None:
            fill, line = PANEL, LINE
        else:
            fill, line = RGBColor(*raw), RED
        box = _card(slide, x, top, width, height, fill=fill, line=line)
        box.name = f"Card {index + 1}"
        lines = [(card["header"], True, 0)]
        lines += [(item, False, 0) for item in card["lines"]]
        _set_lines(box, lines, size=BOX_PT, margin=0.16)


def render_math(slide, spec):
    """Full-width equation rows above a full-width reading panel.

    The V0.23 pilot put prose in a 4.25" column beside the math hosts.  At the
    mandated 20 pt that column holds ~30 characters per line, which is too
    narrow for successor prose, so the deck stacks instead of splitting.
    """
    hosts = spec["hosts"]
    height = 1.00 if len(hosts) == 1 else 0.85
    gap = 0.14
    for index, host in enumerate(hosts):
        y = 1.45 + index * (height + gap)
        _text_box(slide, 0.62, y + (height - 0.52) / 2, 0.86, 0.52,
                  host["label"], size=BOX_PT, bold=True,
                  color=GREEN if host.get("accent") else MUTED)
        _formula_host(slide, host["id"], 1.52, y, 11.19, height,
                      fill=PALE_GREEN if host.get("accent") else WHITE,
                      line=GREEN if host.get("accent") else LINE)
    top = 1.45 + len(hosts) * (height + gap) + 0.10
    bottom = 6.38 if spec.get("note") else 6.90
    panel = _card(slide, 0.62, top, 12.09, bottom - top, fill=PANEL, line=LINE)
    panel.name = "Reading panel"
    lines = [(spec["left_header"], True, 0)]
    lines += [(item, False, 0) for item in spec["left_lines"]]
    _set_lines(panel, lines, size=BOX_PT, margin=0.16)
    if spec.get("note"):
        _footer(slide, spec["note"])


def render_quote(slide, spec):
    box = _card(slide, 0.62, 1.45, 12.09, 2.30, fill=PALE_GREEN, line=GREEN)
    box.name = "Verbatim boundary sentence"
    _set_text(box, spec["quote"], size=BOX_PT, color=INK,
              align=PP_ALIGN.LEFT, margin=0.20)
    below = _card(slide, 0.62, 3.90, 12.09, 3.00, fill=PANEL, line=LINE)
    lines = [(spec["header"], True, 0)]
    lines += [(item, False, 0) for item in spec["lines"]]
    _set_lines(below, lines, size=BOX_PT, margin=0.16)


def render_table(slide, spec):
    headers = spec["headers"]
    rows = spec["rows"]
    top, left, total = 1.45, 0.62, 12.09
    widths = spec.get("widths") or [total / len(headers)] * len(headers)
    header_h = 0.52
    bottom = 6.34 if spec.get("footer") else 7.00
    body_h = (bottom - top - header_h) / max(len(rows), 1)
    x = left
    for index, head in enumerate(headers):
        cell = _card(slide, x, top, widths[index], header_h,
                     fill=DEEP_BLUE, line=DEEP_BLUE, radius=False)
        cell.name = f"Header {index + 1}"
        _set_text(cell, head, size=BOX_PT, bold=True, color=WHITE,
                  align=PP_ALIGN.LEFT, margin=0.08)
        x += widths[index]
    for r_index, row in enumerate(rows):
        y = top + header_h + r_index * body_h
        x = left
        fill = WHITE if r_index % 2 == 0 else PANEL
        for c_index, value in enumerate(row):
            cell = _card(slide, x, y, widths[c_index], body_h,
                         fill=fill, line=LINE, radius=False)
            cell.name = f"Cell {r_index + 1}-{c_index + 1}"
            _set_text(cell, value, size=BOX_PT, color=INK,
                      align=PP_ALIGN.LEFT, margin=0.08)
            x += widths[c_index]


RENDERERS = {
    "title": render_title,
    "cards": render_cards,
    "math": render_math,
    "quote": render_quote,
    "table": render_table,
}


# --------------------------------------------------------------------------
# build
# --------------------------------------------------------------------------

def build(deck) -> Path:
    prs = Presentation(TEMPLATE)
    blank = prs.slide_masters[0].slide_layouts[6]

    for spec in deck:
        slide = prs.slides.add_slide(blank)
        for shape in list(slide.shapes):
            if shape.is_placeholder and shape.placeholder_format.idx == 10:
                continue  # keep the template slide-number field
            _remove_shape(shape)
        if spec["kind"] != "title":
            _chrome(slide, spec["title"], spec.get("badge"), spec.get("subtitle"))
        RENDERERS[spec["kind"]](slide, spec)
        if spec.get("footer"):
            _footer(slide, spec["footer"], tone=spec.get("footer_tone", "amber"))
        _notes(slide, spec["notes"])

    # drop the three template sample slides, keeping master and layouts
    slide_ids = prs.slides._sldIdLst
    for item in list(slide_ids)[:3]:
        slide_ids.remove(item)

    manifest = native_math.load_manifest(MANIFEST)
    injected = native_math.inject(prs, manifest)
    overflow = fit_report(prs)

    prs.save(OUTPUT)
    REPORT.write_text(json.dumps({
        "template": str(TEMPLATE),
        "output": str(OUTPUT),
        "slides": len(deck),
        "native_equations": injected,
        "overflow_estimate": overflow,
        "overflow_note": ("static metric-free estimate at 0.46 em average advance; "
                          "no renderer is installed on this host, so visual fit is "
                          "estimated, not observed"),
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return OUTPUT


def fit_report(prs) -> list[dict]:
    """Estimate text overflow without a renderer.

    Times New Roman averages ~0.46 em of advance over mixed-case English prose.
    This is an estimate, not a measurement: no LibreOffice or PowerPoint is
    installed on this host, so the deck's visual fit is unverified.
    """
    import math

    out = []
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            frame = shape.text_frame
            width = (shape.width - (frame.margin_left or 0)
                     - (frame.margin_right or 0)) / 914400
            height = (shape.height - (frame.margin_top or 0)
                      - (frame.margin_bottom or 0)) / 914400
            paragraphs = [p for p in frame.paragraphs
                          if "".join(r.text for r in p.runs).strip()]
            need = 0.0
            for order, paragraph in enumerate(paragraphs):
                text = "".join(r.text for r in paragraph.runs)
                size = max((r.font.size.pt for r in paragraph.runs
                            if r.font.size), default=BOX_PT)
                rows = max(1, math.ceil(len(text) * size * 0.46 / 72 / width))
                after = (paragraph.space_after.pt if paragraph.space_after else 0) / 72
                need += rows * size * 1.20 / 72
                if order < len(paragraphs) - 1:
                    need += after
            if need > height:
                out.append({"slide": index, "shape": shape.name,
                            "needed_in": round(need, 2),
                            "available_in": round(height, 2),
                            "ratio": round(need / height, 2)})
    return out


if __name__ == "__main__":
    from deck_content import DECK
    print(build(DECK))
